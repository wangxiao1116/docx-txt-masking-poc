from __future__ import annotations

from io import BytesIO
from pathlib import Path

from charset_normalizer import from_bytes
from docx import Document

from models import ClassificationLabel, ReplacementMapping, VerificationResult
from docx_metadata import read_custom_properties


def _extract_docx_text(output_bytes: bytes) -> str:
    doc = Document(BytesIO(output_bytes))
    parts: list[str] = []
    parts.extend(paragraph.text for paragraph in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(paragraph.text for paragraph in cell.paragraphs)
    for section in doc.sections:
        parts.extend(paragraph.text for paragraph in section.header.paragraphs)
        parts.extend(paragraph.text for paragraph in section.footer.paragraphs)
    return "\n".join(parts)


def _extract_txt_text(output_bytes: bytes) -> str:
    detected = from_bytes(output_bytes).best()
    if detected is None:
        raise ValueError("Unable to decode output TXT")
    return str(detected)


def verify_output(
    filename: str,
    output_bytes: bytes,
    label: ClassificationLabel,
    mappings: list[ReplacementMapping],
) -> VerificationResult:
    suffix = Path(filename).suffix.lower()
    details: list[str] = []

    try:
        if suffix == ".docx":
            text = _extract_docx_text(output_bytes)
        elif suffix == ".txt":
            text = _extract_txt_text(output_bytes)
        else:
            raise ValueError("Unsupported format")
        file_openable = True
    except Exception as exc:
        return VerificationResult(
            file_openable=False,
            entity_count=len(mappings),
            original_value_residual_count=-1,
            masked_value_count=0,
            label_present=False,
            passed=False,
            details=[f"Output file cannot be reopened: {exc}"],
        )

    unique_pairs = {(m.original_value, m.masked_value) for m in mappings}
    original_residual = sum(text.count(original) for original, _ in unique_pairs)
    masked_count = sum(text.count(masked) for _, masked in unique_pairs)
    if suffix == ".docx":
        properties = read_custom_properties(output_bytes)
        label_present = (
            properties.get("DataCategoryCode") == label.category_code
            and properties.get("DataCategoryName") == label.category_name
            and properties.get("DataLevel") == label.level
            and properties.get("LabelSource") == label.source
        )
    else:
        label_present = (
            label.category_code in text
            and label.category_name in text
            and label.level in text
        )

    if not mappings:
        details.append("未提取到敏感实体，不能判定脱敏闭环通过")
    if original_residual:
        details.append(f"发现 {original_residual} 处原始敏感值残留")
    else:
        details.append("未发现原始敏感值残留")

    details.append(f"发现 {masked_count} 处遮蔽值")
    details.append("分类分级标签存在" if label_present else "分类分级标签缺失")

    passed = (
        file_openable
        and len(mappings) > 0
        and original_residual == 0
        and masked_count > 0
        and label_present
    )

    return VerificationResult(
        file_openable=file_openable,
        entity_count=len(mappings),
        original_value_residual_count=original_residual,
        masked_value_count=masked_count,
        label_present=label_present,
        passed=passed,
        details=details,
    )
