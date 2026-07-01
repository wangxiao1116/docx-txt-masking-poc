from __future__ import annotations

from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SAMPLE_DIR = ROOT / "sample-data"
SAMPLE_DIR.mkdir(parents=True, exist_ok=True)

demo_id = "110105" + "19491231002X"

(SAMPLE_DIR / "sample_labeled.txt").write_text(
    f"分类标签：A1-4 个人信息\n分级标签：第3级\n员工身份证号码：{demo_id}\n",
    encoding="utf-8",
)

document = Document()
document.add_paragraph(f"分类标签：A1-4 个人信息")
document.add_paragraph("分级标签：第3级")
document.add_paragraph(f"员工身份证号码：{demo_id}")
document.save(SAMPLE_DIR / "sample_labeled.docx")
