from __future__ import annotations

from collections.abc import Iterable
from io import BytesIO
from pathlib import Path

from charset_normalizer import from_bytes
from docx import Document
from docx.document import Document as _Document
from docx.table import _Cell
from docx.text.paragraph import Paragraph


PREVIEW_LIMIT = 12000


def _iter_paragraphs(parent: _Document | _Cell) -> Iterable[Paragraph]:
    for paragraph in parent.paragraphs:
        yield paragraph
    for table in parent.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_paragraphs(cell)


def _docx_text(file_bytes: bytes) -> str:
    document = Document(BytesIO(file_bytes))
    parts: list[str] = []
    parts.extend(paragraph.text for paragraph in _iter_paragraphs(document))
    for section in document.sections:
        parts.extend(paragraph.text for paragraph in _iter_paragraphs(section.header))
        parts.extend(paragraph.text for paragraph in _iter_paragraphs(section.footer))
    return "\n".join(part for part in parts if part)


def _txt_text(file_bytes: bytes) -> str:
    detected = from_bytes(file_bytes).best()
    if detected is None:
        raise ValueError("TXT 编码无法识别")
    return str(detected)


def extract_preview(filename: str, file_bytes: bytes, limit: int = PREVIEW_LIMIT) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".docx":
        text = _docx_text(file_bytes)
    elif suffix == ".txt":
        text = _txt_text(file_bytes)
    else:
        raise ValueError("文件格式不支持预览")
    if len(text) <= limit:
        return text
    return f"{text[:limit]}\n\n[预览已截断，剩余 {len(text) - limit} 个字符未显示]"
