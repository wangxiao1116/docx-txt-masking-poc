from __future__ import annotations

import csv
import sys
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient

API_ROOT = Path(__file__).resolve().parents[1]
if str(API_ROOT) not in sys.path:
    sys.path.insert(0, str(API_ROOT))
if "app" in sys.modules and not hasattr(sys.modules["app"], "__path__"):
    del sys.modules["app"]

from app.db.session import Base, engine
from app.main import app


VALID_ID = "110105" + "19491231002X"


def _docx_bytes(text: str, split_run: bool = False) -> bytes:
    document = Document()
    paragraph = document.add_paragraph()
    if split_run:
        paragraph.add_run(text[:10])
        paragraph.add_run(text[10:])
    else:
        paragraph.add_run(text)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _reset_state() -> None:
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)
    root = Path(__file__).resolve().parents[3] / "storage"
    for folder in ["originals", "outputs", "exports"]:
        target = root / folder
        target.mkdir(parents=True, exist_ok=True)
        for item in target.iterdir():
            if item.name != ".gitkeep" and item.is_file():
                item.unlink()


def _upload(client: TestClient, filename: str, content: bytes):
    return client.post(
        "/api/v1/files",
        data={
            "category_code": "A1-1",
            "category_name": "个人信息",
            "level": "第3级",
            "note": "自动化测试",
        },
        files={"file": (filename, content)},
    )


def test_docx_upload_process_download_and_metadata():
    _reset_state()
    client = TestClient(app)

    response = _upload(client, "sample.docx", _docx_bytes(f"身份证号：{VALID_ID}"))

    assert response.status_code == 200
    payload = response.json()["data"]
    assert payload["file_type"] == "DOCX"
    assert payload["entity_count"] == 1
    assert payload["processed_entity_count"] == 1
    assert payload["process_status"] == "处理成功"
    assert payload["verification_status"] == "通过"

    download = client.get(f"/api/v1/files/{payload['id']}/download")
    assert download.status_code == 200
    reopened = Document(BytesIO(download.content))
    body_text = "\n".join(paragraph.text for paragraph in reopened.paragraphs)
    assert "110105********002X" in body_text
    assert "A1-1" not in body_text
    assert "个人信息" not in body_text


def test_txt_upload_process_and_no_entity_status():
    _reset_state()
    client = TestClient(app)

    success = _upload(client, "sample.txt", f"身份证号：{VALID_ID}".encode("utf-8"))
    clean = _upload(client, "clean.txt", "没有身份证信息".encode("utf-8"))

    assert success.status_code == 200
    assert success.json()["data"]["file_type"] == "TXT"
    assert success.json()["data"]["process_status"] == "处理成功"
    assert clean.status_code == 200
    assert clean.json()["data"]["process_status"] == "未发现敏感实体"
    assert clean.json()["data"]["verification_status"] == "未通过"


def test_file_detail_and_previews_return_200_for_txt_and_docx():
    _reset_state()
    client = TestClient(app)

    txt_task = _upload(client, "detail.txt", f"ID {VALID_ID}".encode("utf-8")).json()["data"]
    docx_task = _upload(client, "detail.docx", _docx_bytes(f"ID {VALID_ID}")).json()["data"]

    for task in [txt_task, docx_task]:
        detail = client.get(f"/api/v1/files/{task['id']}")
        original_preview = client.get(f"/api/v1/files/{task['id']}/preview/original")
        masked_preview = client.get(f"/api/v1/files/{task['id']}/preview/masked")

        assert detail.status_code == 200
        assert original_preview.status_code == 200
        assert masked_preview.status_code == 200


def test_txt_label_is_saved_as_task_metadata_not_file_body():
    _reset_state()
    client = TestClient(app)

    task = _upload(client, "sample.txt", f"身份证号：{VALID_ID}".encode("utf-8")).json()["data"]
    download = client.get(f"/api/v1/files/{task['id']}/download")
    text = download.content.decode("utf-8")

    assert "110105********002X" in text
    assert "category_code=A1-1" not in text
    assert "category_name=个人信息" not in text
    assert "level=第3级" not in text


def test_cross_run_docx_entity_is_masked():
    _reset_state()
    client = TestClient(app)

    response = _upload(client, "split.docx", _docx_bytes(f"身份证号：{VALID_ID}", split_run=True))

    assert response.status_code == 200
    task_id = response.json()["data"]["id"]
    masked_preview = client.get(f"/api/v1/files/{task_id}/preview/masked").json()["data"]["text"]
    assert "110105********002X" in masked_preview
    assert VALID_ID not in masked_preview


def test_manifest_csv_and_results_zip_include_processed_files():
    _reset_state()
    client = TestClient(app)
    _upload(client, "a.docx", _docx_bytes(f"身份证号：{VALID_ID}"))
    _upload(client, "a.txt", f"身份证号：{VALID_ID}".encode("utf-8"))

    csv_response = client.get("/api/v1/exports/manifest.csv")
    assert csv_response.status_code == 200
    csv_text = csv_response.content.decode("utf-8-sig")
    rows = list(csv.DictReader(StringIO(csv_text)))
    assert {row["文件名"] for row in rows} == {"a.docx", "a.txt"}
    assert rows[0]["分类编码"] == "A1-1"

    zip_response = client.get("/api/v1/exports/results.zip")
    assert zip_response.status_code == 200
    with ZipFile(BytesIO(zip_response.content), "r") as archive:
        names = set(archive.namelist())
    assert "处理清单.csv" in names
    assert "脱敏文件/a_masked.docx" in names
    assert "脱敏文件/a_masked.txt" in names


def test_delete_task_removes_record_and_files():
    _reset_state()
    client = TestClient(app)
    task = _upload(client, "sample.txt", f"身份证号：{VALID_ID}".encode("utf-8")).json()["data"]

    response = client.delete(f"/api/v1/files/{task['id']}")

    assert response.status_code == 200
    assert client.get(f"/api/v1/files/{task['id']}").status_code == 404


def test_unsupported_format_is_rejected():
    _reset_state()
    client = TestClient(app)

    response = _upload(client, "sample.pdf", b"not supported")

    assert response.status_code == 400
    assert response.json()["code"] == 40001


def test_database_does_not_store_full_id_card_value():
    _reset_state()
    client = TestClient(app)
    task = _upload(client, "sample.txt", f"身份证号：{VALID_ID}".encode("utf-8")).json()["data"]

    entities = client.get(f"/api/v1/files/{task['id']}/entities").json()["data"]

    assert entities
    assert all(entity["masked_original_value"] != VALID_ID for entity in entities)
    assert all("********" in entity["masked_original_value"] for entity in entities)
