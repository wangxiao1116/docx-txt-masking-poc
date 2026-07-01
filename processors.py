from __future__ import annotations

from collections.abc import Iterable
from io import BytesIO
from pathlib import Path
from typing import Any

from charset_normalizer import from_bytes
from docx import Document
from docx.document import Document as _Document
from docx.table import _Cell, Table
from docx.text.paragraph import Paragraph

from entity_extractor import extract_id_cards
from masking import build_replacement_mappings
from models import ClassificationLabel, ReplacementMapping, SensitiveEntity


def _iter_block_paragraphs(parent: _Document | _Cell) -> Iterable[Paragraph]:
    for paragraph in parent.paragraphs:
        yield paragraph
    for table in parent.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_block_paragraphs(cell)


def _iter_all_docx_paragraphs(doc: _Document) -> Iterable[tuple[str, Paragraph]]:
    for idx, paragraph in enumerate(_iter_block_paragraphs(doc)):
        yield f"body:{idx}", paragraph

    for section_idx, section in enumerate(doc.sections):
        for idx, paragraph in enumerate(_iter_block_paragraphs(section.header)):
            yield f"header:{section_idx}:{idx}", paragraph
        for idx, paragraph in enumerate(_iter_block_paragraphs(section.footer)):
            yield f"footer:{section_idx}:{idx}", paragraph


def _replace_in_paragraph(paragraph: Paragraph, replacements: dict[str, str]) -> int:
    """Replace values across DOCX runs while keeping the existing run objects."""
    if not paragraph.runs or not replacements:
        return 0

    original_run_texts = [run.text for run in paragraph.runs]
    full_text = "".join(original_run_texts)
    if not full_text:
        return 0

    spans: list[tuple[int, int, str, str]] = []
    for original, masked in replacements.items():
        start = 0
        while True:
            idx = full_text.find(original, start)
            if idx < 0:
                break
            spans.append((idx, idx + len(original), original, masked))
            start = idx + len(original)

    if not spans:
        return 0

    # Build character-index to run-index mapping.
    boundaries: list[tuple[int, int]] = []
    cursor = 0
    for text in original_run_texts:
        boundaries.append((cursor, cursor + len(text)))
        cursor += len(text)

    replaced_count = 0
    # Process from right to left so original offsets remain valid.
    for start, end, _original, masked in sorted(spans, key=lambda item: item[0], reverse=True):
        start_run = None
        end_run = None
        for idx, (run_start, run_end) in enumerate(boundaries):
            if start_run is None and run_start <= start < run_end:
                start_run = idx
            if run_start < end <= run_end:
                end_run = idx
                break

        if start_run is None or end_run is None:
            continue

        start_offset = start - boundaries[start_run][0]
        end_offset = end - boundaries[end_run][0]

        if start_run == end_run:
            current = paragraph.runs[start_run].text
            paragraph.runs[start_run].text = (
                current[:start_offset] + masked + current[end_offset:]
            )
        else:
            first_text = paragraph.runs[start_run].text
            last_text = paragraph.runs[end_run].text
            prefix = first_text[:start_offset]
            suffix = last_text[end_offset:]
            paragraph.runs[start_run].text = prefix + masked + suffix
            for idx in range(start_run + 1, end_run + 1):
                paragraph.runs[idx].text = ""

        replaced_count += 1

    return replaced_count


def _docx_text(doc: _Document) -> str:
    return "\n".join(paragraph.text for _, paragraph in _iter_all_docx_paragraphs(doc))


def process_docx(
    source_bytes: bytes,
    label: ClassificationLabel,
) -> tuple[bytes, list[SensitiveEntity], list[ReplacementMapping], dict[str, Any]]:
    doc = Document(BytesIO(source_bytes))

    entities: list[SensitiveEntity] = []
    for location, paragraph in _iter_all_docx_paragraphs(doc):
        entities.extend(extract_id_cards(paragraph.text, location))

    mappings = build_replacement_mappings(entities)
    replacement_dict = {
        mapping.original_value: mapping.masked_value for mapping in mappings
    }

    replacement_count = 0
    for _location, paragraph in _iter_all_docx_paragraphs(doc):
        replacement_count += _replace_in_paragraph(paragraph, replacement_dict)

    output = BytesIO()
    doc.save(output)

    from docx_metadata import write_custom_properties

    label_properties = {
        "DataCategoryCode": label.category_code,
        "DataCategoryName": label.category_name,
        "DataLevel": label.level,
        "LabelSource": label.source,
    }
    output_bytes = write_custom_properties(output.getvalue(), label_properties)

    metadata = {
        "format": "DOCX",
        "entity_count": len(entities),
        "replacement_count": replacement_count,
        "label_storage": "DOCX_CUSTOM_PROPERTIES",
        "label_properties": label_properties,
    }
    return output_bytes, entities, mappings, metadata


def _decode_txt(source_bytes: bytes) -> tuple[str, str]:
    detected = from_bytes(source_bytes).best()
    if detected is None:
        raise ValueError("Unable to detect TXT encoding")
    return str(detected), detected.encoding or "utf-8"


def process_txt(
    source_bytes: bytes,
    label: ClassificationLabel,
) -> tuple[bytes, list[SensitiveEntity], list[ReplacementMapping], dict[str, Any]]:
    text, encoding = _decode_txt(source_bytes)
    entities = extract_id_cards(text, "txt:full")
    mappings = build_replacement_mappings(entities)

    output_text = text
    for mapping in mappings:
        output_text = output_text.replace(mapping.original_value, mapping.masked_value)

    label_header = "\n".join(
        [
            "# DATASET_METADATA_BEGIN",
            f"# category_code={label.category_code}",
            f"# category_name={label.category_name}",
            f"# level={label.level}",
            f"# label_source={label.source}",
            "# DATASET_METADATA_END",
            "",
        ]
    )
    output_text = label_header + output_text

    metadata = {
        "format": "TXT",
        "encoding": encoding,
        "entity_count": len(entities),
        "replacement_count": len(mappings),
        "label_text": label_header,
    }
    return output_text.encode(encoding, errors="strict"), entities, mappings, metadata


def process_file(
    filename: str,
    source_bytes: bytes,
    label: ClassificationLabel,
) -> tuple[bytes, list[SensitiveEntity], list[ReplacementMapping], dict[str, Any]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".docx":
        return process_docx(source_bytes, label)
    if suffix == ".txt":
        return process_txt(source_bytes, label)
    raise ValueError("Only .docx and .txt files are supported in this POC")
