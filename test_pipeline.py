from io import BytesIO

from docx import Document

from models import ClassificationLabel
from processors import process_docx, process_txt
from verification import verify_output
from docx_metadata import read_custom_properties


VALID_ID = "11010519491231002X"
SYNTHETIC_ID = "21010219950820123X"
LABEL = ClassificationLabel("A1-1", "自然人身份标识", "第3级")


def test_txt_pipeline():
    source = f"公民身份号码：{VALID_ID}\n".encode("utf-8")
    output, entities, mappings, _ = process_txt(source, LABEL)
    verification = verify_output("sample.txt", output, LABEL, mappings)

    assert len(entities) == 1
    assert entities[0].checksum_valid is True
    assert "110105********002X" in output.decode("utf-8")
    assert verification.passed


def test_docx_pipeline_cross_run():
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("公民身份号码：110105")
    paragraph.add_run("19491231002X")
    source = BytesIO()
    document.save(source)

    output, entities, mappings, _ = process_docx(source.getvalue(), LABEL)
    verification = verify_output("sample.docx", output, LABEL, mappings)

    assert len(entities) == 1
    properties = read_custom_properties(output)
    assert properties["DataCategoryCode"] == LABEL.category_code
    assert properties["DataCategoryName"] == LABEL.category_name
    assert properties["DataLevel"] == LABEL.level
    reopened = Document(BytesIO(output))
    body_text = "\n".join(p.text for p in reopened.paragraphs)
    assert "【数据分类标签】" not in body_text
    assert verification.passed


def test_synthetic_invalid_checksum_is_still_masked():
    document = Document()
    document.add_paragraph(f"公民身份号码：{SYNTHETIC_ID}")
    source = BytesIO()
    document.save(source)

    output, entities, mappings, _ = process_docx(source.getvalue(), LABEL)
    verification = verify_output("sample.docx", output, LABEL, mappings)

    assert len(entities) == 1
    assert entities[0].checksum_valid is False
    assert verification.passed


def test_no_entity_does_not_pass():
    source = "无身份证信息".encode("utf-8")
    output, entities, mappings, _ = process_txt(source, LABEL)
    verification = verify_output("sample.txt", output, LABEL, mappings)

    assert len(entities) == 0
    assert verification.passed is False
