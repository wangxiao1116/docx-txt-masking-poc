from __future__ import annotations

from collections.abc import Iterable
from io import BytesIO
from pathlib import Path

from charset_normalizer import from_bytes
from docx import Document
from docx.document import Document as _Document
from docx.table import _Cell
from docx.text.paragraph import Paragraph


DEFAULT_PREVIEW_LIMIT = 8000


def _iter_paragraphs(parent: _Document | _Cell) -> Iterable[Paragraph]:
    for paragraph in parent.paragraphs:
        yield paragraph
    for table in parent.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_paragraphs(cell)


def _extract_docx_text(file_bytes: bytes) -> str:
    document = Document(BytesIO(file_bytes))
    parts: list[str] = []

    parts.extend(paragraph.text for paragraph in _iter_paragraphs(document))

    for section in document.sections:
        parts.extend(paragraph.text for paragraph in _iter_paragraphs(section.header))
        parts.extend(paragraph.text for paragraph in _iter_paragraphs(section.footer))

    return "\n".join(part for part in parts if part)


def _extract_txt_text(file_bytes: bytes) -> str:
    detected = from_bytes(file_bytes).best()
    if detected is None:
        raise ValueError("Unable to decode TXT preview")
    return str(detected)


def truncate_preview(text: str, max_chars: int = DEFAULT_PREVIEW_LIMIT) -> str:
    if max_chars <= 0 or len(text) <= max_chars:
        return text
    omitted = len(text) - max_chars
    return f"{text[:max_chars]}\n\n[预览已截断，剩余 {omitted} 个字符未显示]"


def extract_preview_text(
    filename: str,
    file_bytes: bytes,
    max_chars: int = DEFAULT_PREVIEW_LIMIT,
) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".docx":
        text = _extract_docx_text(file_bytes)
    elif suffix == ".txt":
        text = _extract_txt_text(file_bytes)
    else:
        raise ValueError("Only .docx and .txt previews are supported")
    return truncate_preview(text, max_chars=max_chars)
