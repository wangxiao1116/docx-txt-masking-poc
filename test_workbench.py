from __future__ import annotations

from io import BytesIO
from zipfile import ZipFile

from docx import Document

from docx_metadata import read_custom_properties
from models import ClassificationLabel
from processors import process_docx, process_txt
from record_manager import create_file_record
from export_service import generate_manifest_csv, generate_results_zip
from verification import verify_output


VALID_ID = "11010519491231002X"
LABEL = ClassificationLabel("A1-1", "个人信息", "第3级")


def _make_docx_bytes(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _process_record(filename: str, source_bytes: bytes):
    if filename.endswith(".docx"):
        output, entities, mappings, metadata = process_docx(source_bytes, LABEL)
    else:
        output, entities, mappings, metadata = process_txt(source_bytes, LABEL)
    verification = verify_output(filename, output, LABEL, mappings)
    return create_file_record(
        filename=filename,
        source_bytes=source_bytes,
        output_bytes=output,
        label=LABEL,
        entities=entities,
        mappings=mappings,
        verification=verification,
        metadata=metadata,
        remark="test remark",
    )


def test_docx_and_txt_can_form_two_records():
    docx_record = _process_record("sample.docx", _make_docx_bytes(f"身份证号：{VALID_ID}"))
    txt_record = _process_record(
        "sample.txt",
        f"身份证号：{VALID_ID}".encode("utf-8"),
    )

    assert len([docx_record, txt_record]) == 2
    assert docx_record.file_format == "DOCX"
    assert txt_record.file_format == "TXT"
    assert docx_record.processing_status == "处理成功"
    assert txt_record.processing_status == "处理成功"


def test_csv_contains_required_manifest_fields():
    records = [
        _process_record("sample.docx", _make_docx_bytes(f"身份证号：{VALID_ID}")),
        _process_record("sample.txt", f"身份证号：{VALID_ID}".encode("utf-8")),
    ]

    csv_bytes = generate_manifest_csv(records)
    csv_text = csv_bytes.decode("utf-8-sig")

    assert csv_bytes.startswith(b"\xef\xbb\xbf")
    assert "文件名" in csv_text
    assert "分类编码" in csv_text
    assert "分类名称" in csv_text
    assert "分级标签" in csv_text
    assert "敏感实体数量" in csv_text
    assert "处理敏感实体数量" in csv_text
    assert "sample.docx" in csv_text
    assert "A1-1" in csv_text


def test_zip_contains_manifest_and_all_masked_files():
    records = [
        _process_record("sample.docx", _make_docx_bytes(f"身份证号：{VALID_ID}")),
        _process_record("sample.txt", f"身份证号：{VALID_ID}".encode("utf-8")),
    ]

    zip_bytes = generate_results_zip(records)

    with ZipFile(BytesIO(zip_bytes), "r") as archive:
        names = set(archive.namelist())

    assert "处理清单.csv" in names
    assert "sample_masked.docx" in names
    assert "sample_masked.txt" in names


def test_no_sensitive_entity_status_and_verification_do_not_pass():
    record = _process_record("clean.txt", "没有身份证信息".encode("utf-8"))

    assert record.sensitive_entity_count == 0
    assert record.processing_status == "未发现敏感实体"
    assert record.verification_status == "未通过"
    assert record.verification.passed is False


def test_docx_label_still_uses_custom_properties_not_body_text():
    source_bytes = _make_docx_bytes(f"身份证号：{VALID_ID}")
    record = _process_record("sample.docx", source_bytes)

    properties = read_custom_properties(record.output_bytes)
    reopened = Document(BytesIO(record.output_bytes))
    body_text = "\n".join(paragraph.text for paragraph in reopened.paragraphs)

    assert properties["DataCategoryCode"] == LABEL.category_code
    assert properties["DataCategoryName"] == LABEL.category_name
    assert properties["DataLevel"] == LABEL.level
    assert LABEL.category_code not in body_text
    assert LABEL.category_name not in body_text
    assert LABEL.level not in body_text
